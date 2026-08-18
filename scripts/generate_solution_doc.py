from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "Growth_OS_笔试方案.docx"


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    props.append(element)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)
    return paragraph


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, "315B3F")
        set_cell_text(cell, value, bold=True, color="FFFFFF")
        if widths:
            cell.width = Cm(widths[index])
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            if widths:
                cells[index].width = Cm(widths[index])
    doc.add_paragraph()
    return table


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.35

    for style_name, size, color in [("Title", 28, "315B3F"), ("Heading 1", 18, "315B3F"), ("Heading 2", 13, "2F4C38")]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Growth OS | AI Content Operations | 笔试方案")
    set_run_font(footer_run, 8, color="849089")


def add_cover(doc):
    doc.add_paragraph()
    eyebrow = doc.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = eyebrow.add_run("AI CONTENT OPERATIONS · WRITTEN ASSIGNMENT DEMO")
    set_run_font(run, 10, bold=True, color="5B936E")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Growth OS")
    set_run_font(run, 32, bold=True, color="315B3F")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("AI 驱动的内容增长运营工作台方案")
    set_run_font(run, 17, color="425D49")

    doc.add_paragraph()
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = line.add_run("从实时信号发现，到内容生产、分发决策与复盘沉淀的可运行闭环")
    set_run_font(run, 11, color="6E7B70")

    doc.add_paragraph()
    card = doc.add_table(rows=3, cols=2)
    card.style = "Light Shading Accent 1"
    values = [("方案定位", "面向多账号内容运营团队的 AI 增长工作台"), ("交付形态", "React + Node.js 可运行 Demo，支持本地状态持久化"), ("覆盖范围", "笔试任务 1-5，重点将任务 2/3/4 整合为增长情报闭环")]
    for row, (label, value) in zip(card.rows, values):
        shade(row.cells[0], "E8F1E8")
        set_cell_text(row.cells[0], label, bold=True, color="315B3F")
        set_cell_text(row.cells[1], value)
    doc.add_page_break()


def main():
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "一、方案概述", 1)
    doc.add_paragraph("Growth OS 是一个面向内容运营团队的 AI 工作台。它将原本分散的热点发现、账号复盘、对标研究、爆款拆解和发布前优化，连接成一条可追踪、可人工决策、可沉淀的数据闭环。")
    doc.add_paragraph("方案并不把 AI 设计成“自动替人发内容”的黑盒，而是把 AI 的价值放在信息筛选、结构化判断与动作建议上；运营与 Leader 始终保留内容取舍、审核和发布的最终决策权。")

    add_heading(doc, "1.1 目标用户与核心问题", 2)
    add_table(doc, ["角色", "当前痛点", "Growth OS 的解决方式"], [
        ["内容运营", "热点多、筛选成本高；账号问题靠经验判断", "以热点候选池、账号问题卡和明确下一步动作减少信息噪声"],
        ["运营 Leader", "难以同时掌握多账号风险、任务进度和策略质量", "以总仪表盘、经营侧栏、Leader 任务和日报汇总管理经营动作"],
        ["内容创作者 / 审核人", "容易照搬对标内容，且发布前缺少统一检查", "以机制拆解、表达相似度门槛与人工审核队列降低同质化和风险"],
    ], [2.7, 6.0, 7.0])

    add_heading(doc, "1.2 核心设计原则", 2)
    for item in [
        "闭环优先：每一个 Agent 输出都必须能进入下一步，而不是停留在一份报告里。",
        "人机协同：AI 提供判断依据、证据和建议，人工决定是否采用与何时发布。",
        "多账号经营：既能在总仪表盘看组合表现，也能下钻到单账号具体问题。",
        "可演示、可扩展：无 Token 时仍可完整跑通本地事件流；接入 X API 和模型接口后可替换为真实数据源。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "二、整体业务闭环", 1)
    doc.add_paragraph("Growth OS 将五个笔试任务组织为一条统一链路，而不是五个孤立页面：")
    flow = doc.add_paragraph()
    flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = flow.add_run("实时信号 / 对标内容  →  增长情报判断  →  候选内容池  →  爆款拆解与二创  →  分发优化  →  审核 / 发布 / 复盘")
    set_run_font(run, 12, bold=True, color="315B3F")
    doc.add_paragraph("其中“增长情报”将任务 2、3、4 作为同一个运营决策单元：热点告诉团队哪些机会值得关注；账号复盘告诉团队当前缺什么；对标学习再输出与账号定位匹配、且表达风险可控的策略。三者共同进入候选内容池。")

    add_heading(doc, "三、信息架构与页面方案", 1)
    add_table(doc, ["模块", "核心用途", "面向角色", "关键产出"], [
        ["总仪表盘", "快速识别整体经营状态和优先级", "Leader / 运营负责人", "组合指标、账号健康分布、风险、待办、内容漏斗"],
        ["增长情报", "整合热点、复盘与对标策略", "内容运营 / Leader", "候选内容、账号问题、Leader 任务、日报、差异化策略"],
        ["爆款拆解", "将高表现内容转为可复用的原创方案", "内容运营 / 创作者", "结构拆解、多个创意方向、可编辑英文草稿"],
        ["分发优化", "发布前做经营判断和风险控制", "运营 / 审核人", "发布时间、Hook、CTA、风险建议、审核队列"],
    ], [2.3, 4.4, 3.1, 5.9])

    add_heading(doc, "3.1 总仪表盘：经营驾驶舱", 2)
    for item in [
        "默认首页：展示管理账号数、总曝光、平均点赞率、平均互动率和待处理事项。",
        "账号管理看板：横向比较曝光、点赞率、互动率、近 7 日发帖和连接状态；点击可下钻至对应账号的增长情报驾驶舱。",
        "经营辅助栏：集中展示高优风险、Leader 待办、审核队列和近期活动；支持折叠并记住用户偏好。",
        "连接状态为 Demo：已连接、待授权、Token 即将过期、数据异常，用于说明未来账号授权管理形态，不发起真实 OAuth。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3.2 增长情报：任务 2/3/4 的统一工作台", 2)
    add_table(doc, ["子能力", "输入", "AI / 系统处理", "人工动作"], [
        ["热点候选池 Agent", "本地实时事件流、未来可接 X API、手动导入", "按动量、相关度、竞争度生成机会判断和理由", "收录、忽略或进入爆款拆解"],
        ["多账号复盘 + Leader 指导", "账号指标、内容信号、历史任务", "识别 1-3 个问题，输出证据、影响、建议", "将首要问题转为负责人、截止日、验证指标明确的任务"],
        ["对标学习 + 差异化策略", "对标账号内容机制与表现", "区分可借鉴机制与不可复制表达，计算表达相似度", "低风险策略进入候选池；高风险策略必须人工改写"],
    ], [3.0, 3.2, 5.1, 4.4])

    add_heading(doc, "3.3 爆款拆解与分发优化", 2)
    doc.add_paragraph("爆款拆解不复制原文，而是从 Hook、内容结构、情绪、CTA 四个维度提取机制；随后生成多个可选择的原创方向和可编辑英文草稿。")
    doc.add_paragraph("分发优化是内容发布前的最后一道经营判断：它检查开头吸引力、推荐发布时间、互动引导和表达风险。运营逐条确认建议后，把内容送入审核队列；审核通过后再由账号侧正式发布。")

    add_heading(doc, "四、笔试五项任务覆盖说明", 1)
    add_table(doc, ["任务", "在产品中的落点", "已演示能力"], [
        ["任务 1：分发优化", "分发优化模块", "发布前 Hook / 时机 / CTA / 风险建议；动作完成记录；进入人工审核队列"],
        ["任务 2：热点与候选池", "增长情报 - 热点候选池", "实时本地信号、匹配度 / 竞争度筛选、候选池状态管理、进入二创"],
        ["任务 3：多账号复盘与指导", "增长情报 - 多账号复盘", "账号问题、证据、影响、建议、Leader 任务、任务完成与每日汇总"],
        ["任务 4：对标与差异化", "增长情报 - 对标策略", "可学习机制 / 不可照搬表达、相似度门槛、低风险策略进入候选池"],
        ["任务 5：爆款二创", "爆款拆解模块", "爆款结构拆解、多个内容方向、可编辑英文草稿、保存为内容资产"],
    ], [3.1, 4.4, 8.2])

    add_heading(doc, "五、核心交互与状态流转", 1)
    add_table(doc, ["节点", "状态", "流转规则"], [
        ["热点信号", "新发现 / 值得测试 / 已忽略", "运营判断是否收录为候选内容"],
        ["候选内容", "待分析 / 已进入拆解 / 已忽略", "热点与低风险对标策略汇合，避免重复建池"],
        ["Leader 任务", "待执行 / 已完成", "问题转任务后，保留负责人、截止日与验证指标"],
        ["内容资产", "已保存 / 待审核", "二创草稿保存后进入资产库，待人工审核"],
        ["分发建议", "建议采用 / 已完成", "发布前建议由人工逐项确认，记录操作动作"],
    ], [3.0, 3.2, 9.5])

    add_heading(doc, "六、技术实现与可运行性", 1)
    add_table(doc, ["层级", "实现", "目的"], [
        ["前端", "React + Vite", "多模块工作台、数据看板、状态筛选、账号下钻、响应式布局"],
        ["后端", "Node.js 原生 HTTP 服务", "提供监测、候选池、资产、任务、日报和策略晋升接口"],
        ["状态", "JSON 文件持久化", "Demo 重启后仍保留候选内容、任务、审核和活动记录"],
        ["监测", "本地事件流 + 可选 X API v2", "无凭据可稳定演示；配置 Bearer Token 后可使用 recent search"],
        ["内容生成", "本地规则引擎 + 可选 OpenAI 兼容接口", "无模型 Token 可跑通；配置后可调用模型并校验结构化输出"],
        ["质量", "Node Test Runner + Vite Build", "验证增长状态流转、日报、相似度门槛和账号组合指标"],
    ], [2.3, 5.1, 8.3])

    add_heading(doc, "6.1 当前运行方式", 2)
    doc.add_paragraph("本地启动命令：pnpm dev:full。前端地址为 http://127.0.0.1:4173/，后端 API 地址为 http://127.0.0.1:8787/。项目已通过 5 项自动化测试与生产构建验证。")

    add_heading(doc, "七、Demo 演示建议（约 5-8 分钟）", 1)
    steps = [
        "从总仪表盘开始：说明多账号经营指标、内容漏斗和经营侧栏如何帮助 Leader 确定今天优先级。",
        "点击风险账号：下钻到增长情报，展示该账号的具体问题、证据、影响和建议；把问题转为 Leader 任务。",
        "回到热点候选池：选择一个相关度高、竞争较低的实时信号，收录到候选内容池。",
        "查看对标策略：说明系统学习的是内容机制而不是原句；高表达相似度方案被拦截，低风险方案可进入候选池。",
        "进入爆款拆解：展示 Hook / 结构 / 情绪 / CTA 拆解，并选取一条英文草稿保存为内容资产。",
        "进入分发优化：确认发布时间、互动引导与表达风险建议，最后送入审核队列，完成从发现到待发布的闭环。",
    ]
    for index, step in enumerate(steps, 1):
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{index:02d}  ").bold = True
        paragraph.add_run(step)

    add_heading(doc, "八、真实能力边界与下一步迭代", 1)
    doc.add_paragraph("为保证笔试 Demo 的可运行性，当前项目明确区分了“已可运行能力”和“未来接入能力”，不把模拟状态描述为真实生产能力。")
    add_table(doc, ["当前已实现", "下一阶段建议"], [
        ["本地实时事件流、状态持久化、候选池、任务与日报、相似度门槛、内容资产和审核队列", "接入真实 X OAuth 与多账号 Token 安全存储；增加授权回调、权限校验与数据同步任务"],
        ["可选 X recent-search 接口与 OpenAI 兼容模型接口", "接入更多平台数据源与人工导入；沉淀账号级历史内容表现"],
        ["发布前分发建议和操作记录", "补齐发布后指标回流、Top 3 历史内容对比、建议效果归因与下一轮策略调整"],
        ["2 个外部对标账号的机制学习演示", "扩展真实对标库、内容语料、表达相似度模型与审批策略"],
    ], [8.5, 8.5])

    add_heading(doc, "九、方案价值总结", 1)
    doc.add_paragraph("Growth OS 的核心不是多做几个 AI 页面，而是把内容运营中最容易断开的环节连接起来：信号发现有去向、账号问题有负责人、对标学习有边界、内容二创可追溯、发布决策可审核、后续可以回到数据继续优化。")
    conclusion = doc.add_paragraph()
    run = conclusion.add_run("一句话概括：让内容团队从“凭经验追热点”，升级为“用数据和 Agent 共同经营内容增长”。")
    set_run_font(run, 12, bold=True, color="315B3F")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
